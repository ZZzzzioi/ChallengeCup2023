# encoding=utf8
import pandas as pd
import docx
from sklearn.linear_model import LogisticRegression as lr
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler

# import data and judge whether it is a experimental unit
data = pd.read_excel('/Users/zzzzzioi_/Library/CloudStorage/OneDrive-个人/研究生打工/2023.03/挑战杯/现状指数-云雨图.xlsx')
experimental_unit_docx = docx.Document(r'/Users/zzzzzioi_/Library/CloudStorage/OneDrive-个人/研究生打工/2023.03/挑战杯/试点名单.docx')
experimental_unit_tables = experimental_unit_docx.tables[0]
exp_unit_list = [[experimental_unit_tables.cell(i, 0).text, experimental_unit_tables.cell(i, 1).text] for i in range(1, len(experimental_unit_tables.rows))]
data_district = [data.iloc[i, 0].split('-') for i in range(len(data))]
data['if_exp_unit'] = None
for unit in range(len(data_district)):
    for exp_list in exp_unit_list:
        if data_district[unit][0] in exp_list[0]:
            if data_district[unit][-1] in exp_list[1]:
                data['if_exp_unit'][unit] = 1
                break
            else:
                data['if_exp_unit'][unit] = 0
                break
        else:
            continue

data.to_csv('/Users/zzzzzioi_/Library/CloudStorage/OneDrive-个人/Code/PSM_python/data/PSM_data.csv', index=False)

# add other data from original data sheet
data = pd.read_csv('/Users/zzzzzioi_/Library/CloudStorage/OneDrive-个人/Code/PSM_python/data/PSM_data.csv')
data = data.drop([153])  # 删除安徽省广德市洪桥村
# 直接在excel上进行修改，防止出现编码错误问题
# 修改控制变量表内变量名称、多余空行
data_original = pd.read_excel('/Users/zzzzzioi_/Library/CloudStorage/OneDrive-个人/Code/PSM_python/data/控制变量表.xlsx')
# 修改位置参数，东北=1，华北=2，华东=3，华南=4，华中=5，西北=6，西南=7
data['分组'] = data['分组'].replace(['东北', '华北', '华东', '华南', '华中', '西北', '西南'], [1, 2, 3, 4, 5, 6, 7])
# 按照村名匹配各项数值
# 匹配常住人口总数
data_population = data_original[['调研地', '村名', '村内常住人口总数']]
data = pd.merge(left=data, right=data_population, how='left', on=['调研地', '村名'])

# 匹配是否为贫困村
data_poor = data_original[['调研地', '村名', '是否为贫困村']]
data_poor['是否为贫困村'] = data_poor['是否为贫困村'].replace('是', 1).replace('否', 0)
data = pd.merge(left=data, right=data_poor, how='left', on=['调研地', '村名'])

# 匹配村人均收入
data_income = data_original[['调研地', '村名', '村人均收入']]
data = pd.merge(left=data, right=data_income, how='left', on=['调研地', '村名'])

# 匹配领取低保补助人数
data_dibao = data_original[['调研地', '村名', '领取低保补助人数']]
data = pd.merge(left=data, right=data_dibao, how='left', on=['调研地', '村名'])

# 匹配教育结构
data_education = data_original[['调研地', '村名', '高中、中专', '大专及以上']]
data_education['教育结构'] = (data_education['高中、中专'] + data_education['大专及以上'])/100
data = pd.merge(left=data, right=data_education[['调研地', '村名', '教育结构']], how='left', on=['调研地', '村名'])

# 匹配性别结构
data_gender = data_original[['调研地', '村名', '男']]
data_gender['性别结构'] = data_gender['男']/100
data = pd.merge(left=data, right=data_gender[['调研地', '村名', '性别结构']], how='left', on=['调研地', '村名'])

# 匹配年龄结构
data_age = data_original[['调研地', '村名', '16-40岁']]
data_age['年龄结构'] = data_age['16-40岁']/100
data = pd.merge(left=data, right=data_age[['调研地', '村名', '年龄结构']], how='left', on=['调研地', '村名'])

# 匹配是否对光伏政策进行了推广普及
data_promotion = data_original[['调研地', '村名', '是否对光伏政策进行了推广普及']].replace('是', 1).replace('否', 0)
data = pd.merge(left=data, right=data_promotion[['调研地', '村名', '是否对光伏政策进行了推广普及']], how='left', on=['调研地', '村名'])

# 匹配村干部是否包含大学生村官或选调生
data_officials = data_original[['调研地', '村名', '村干部是否包含大学生村官或选调生']]
data_officials['村干部是否包含大学生村官或选调生'] = data_officials['村干部是否包含大学生村官或选调生'].replace('是', 1).replace('否', 0)
data = pd.merge(left=data, right=data_officials[['调研地', '村名', '村干部是否包含大学生村官或选调生']], how='left', on=['调研地', '村名'])

# 匹配全村每年电力消耗
data_consume = data_original[['调研地', '村名', '全村每年电力消耗']]
data = pd.merge(left=data, right=data_consume[['调研地', '村名', '全村每年电力消耗']], how='left', on=['调研地', '村名'])

# 匹配村内电价
data_price = data_original[['调研地', '村名', '村内电价']]
data = pd.merge(left=data, right=data_price[['调研地', '村名', '村内电价']], how='left', on=['调研地', '村名'])

# 匹配村内电力支出主要用途
data_family = data_original[['调研地', '村名', '1=家庭用电']]
data_family['村内家庭用电所占比例'] = data_family['1=家庭用电']/100
data = pd.merge(left=data, right=data_family[['调研地', '村名', '村内家庭用电所占比例']], how='left', on=['调研地', '村名'])

# 匹配是否出现过电力负荷过高
data_overload = data_original[['调研地', '村名', '是否出现过电力负荷过高']]
data_overload['是否出现过电力负荷过高'] = data_overload['是否出现过电力负荷过高'].replace('是', 1).replace('否', 0)
data = pd.merge(left=data, right=data_overload[['调研地', '村名', '是否出现过电力负荷过高']], how='left', on=['调研地', '村名'])

# 匹配村内是否有除光伏外的其他新能源建设
data_newenergy = data_original[['调研地', '村名', '村内是否有除光伏外的其他新能源建设']]
data_newenergy['村内是否有除光伏外的其他新能源建设'] = data_newenergy['村内是否有除光伏外的其他新能源建设'].replace('是', 1).replace('否', 0)
data = pd.merge(left=data, right=data_newenergy[['调研地', '村名', '村内是否有除光伏外的其他新能源建设']], how='left', on=['调研地', '村名'])
data.to_csv('/Users/zzzzzioi_/Library/CloudStorage/OneDrive-个人/Code/PSM_python/data/PSM_data_reset.csv', index=False, encoding='UTF-8')
# 使用number将数据转存为excel文件，解决数字乱码问题


path = '/Users/zzzzzioi_/Library/CloudStorage/OneDrive-个人/Code/PSM_python/data/PSM_data_reset_no-.xlsx'
df_data = pd.read_excel(path)
T = data.if_exp_unit
X = df_data.loc[:,df_data.columns !='if_exp_unit']
pipe = Pipeline([
    ('scaler', StandardScaler()),
    ('logistic_classifier', lr())
])

pipe.fit(X_encoded, T)
